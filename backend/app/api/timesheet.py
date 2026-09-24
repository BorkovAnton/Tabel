import io
from datetime import date, datetime, time, timedelta
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.orm import Session
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from urllib.parse import quote

from app.database import get_db
from app.models.employee import Employee
from app.models.turnstile_event import TurnstileEvent
from app.models.work_schedule import WorkSchedule
from app.models.timesheet_record import TimesheetRecord
from app.models.department import Department
from app.schemas.timesheet import (
    TimesheetRecordResponse,
    TimesheetCalculateRequest,
    TimesheetCalculateResponse,
    TimesheetUpdateRequest,
    TimesheetReportResponse,
)

def get_all_department_ids(department_id: int, db: Session) -> List[int]:
    """Рекурсивно получает ID подразделения и всех его дочерних подразделений"""
    from app.models.department import Department
    
    result = [department_id]
    children = db.query(Department).filter(Department.parent_id == department_id).all()
    
    for child in children:
        result.extend(get_all_department_ids(child.id, db))
    
    return result

router = APIRouter(prefix="/api/timesheet", tags=["Timesheet"])

# Если время между входом и выходом больше этого значения (в часах) —
# считаем, что сотрудник забыл сделать отметку и смена помечается для проверки.
MAX_NORMAL_SHIFT_HOURS = 13

DEFAULT_LUNCH_MINUTES = 60
DEFAULT_NORM_HOURS = 8.0


# ---------------------------------------------------------------------------
# Вспомогательные функции
# ---------------------------------------------------------------------------

def parse_date(value: str, field_name: str) -> date:
    """Парсит дату в формате YYYY-MM-DD, иначе бросает HTTPException(400)."""
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        raise HTTPException(
            status_code=400,
            detail=f"Неверный формат {field_name}. Ожидается YYYY-MM-DD",
        )


def daterange(start: date, end: date):
    """Генератор дней от start до end включительно."""
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def schedule_lunch_minutes(schedule: Optional[WorkSchedule], day: date) -> int:
    """Обеденный перерыв по графику для конкретного дня недели."""
    if not schedule:
        return DEFAULT_LUNCH_MINUTES
    
    day_of_week = day.weekday()  # 0=Пн, 6=Вс
    
    day_schedule = next(
        (d for d in schedule.days if d.day_of_week == day_of_week),
        None
    )
    
    if day_schedule and day_schedule.is_day_off:
        return 0
    
    if day_schedule:
        return day_schedule.lunch_minutes
    
    return DEFAULT_LUNCH_MINUTES


def schedule_norm_hours(schedule: Optional[WorkSchedule], day: date) -> float:
    """Норма часов за день по графику."""
    if not schedule:
        return DEFAULT_NORM_HOURS
    
    day_of_week = day.weekday()
    
    day_schedule = next(
        (d for d in schedule.days if d.day_of_week == day_of_week),
        None
    )
    
    if not day_schedule or day_schedule.is_day_off:
        return 0.0
    
    if not day_schedule.start_time or not day_schedule.end_time:
        return 0.0
    
    from datetime import datetime, timedelta
    start_dt = datetime.combine(day, day_schedule.start_time)
    end_dt = datetime.combine(day, day_schedule.end_time)
    
    if end_dt <= start_dt:
        end_dt += timedelta(days=1)
    
    gross = (end_dt - start_dt).total_seconds() / 3600
    net = gross - day_schedule.lunch_minutes / 60
    return round(max(net, 0.0), 2)


def calculate_day_record(events: List[TurnstileEvent], lunch_minutes: int):
    """Вычисляет first_in, last_out, fact_hours, needs_review, review_reason
    по списку событий проходной за один день."""

    in_events = [e for e in events if e.event_type == "in"]
    out_events = [e for e in events if e.event_type == "out"]

    first_in = min((e.datetime for e in in_events), default=None)
    last_out = max((e.datetime for e in out_events), default=None)

    fact_hours = 0.0
    needs_review = False
    review_reason = None

    if first_in and not last_out:
        needs_review = True
        review_reason = "Нет отметки выхода"
    elif last_out and not first_in:
        needs_review = True
        review_reason = "Нет отметки входа"
    elif first_in and last_out:
        if last_out < first_in:
            # Смена перешла через полночь
            worked_seconds = (last_out + timedelta(days=1) - first_in).total_seconds()
        else:
            worked_seconds = (last_out - first_in).total_seconds()

        fact_hours = worked_seconds / 3600

        if fact_hours > MAX_NORMAL_SHIFT_HOURS:
            needs_review = True
            review_reason = (
                f"Слишком длинная смена: {fact_hours:.2f} ч. "
                f"(возможно, пропущена отметка)"
            )
        else:
            needs_review = False
            review_reason = None

        fact_hours -= lunch_minutes / 60
        fact_hours = round(fact_hours, 2)
        if fact_hours < 0:
            fact_hours = 0.0

    return first_in, last_out, fact_hours, needs_review, review_reason


# ---------------------------------------------------------------------------
# Эндпоинты
# ---------------------------------------------------------------------------

@router.post("/calculate", response_model=TimesheetCalculateResponse)
def calculate_timesheet(
    payload: TimesheetCalculateRequest,
    db: Session = Depends(get_db),
):
    """Рассчитывает табель на основе событий проходной за указанный период."""

    date_from = parse_date(payload.date_from, "date_from")
    date_to = parse_date(payload.date_to, "date_to")

    if date_from > date_to:
        raise HTTPException(
            status_code=400,
            detail="date_from не может быть позже date_to",
        )

    # Определяем список сотрудников для расчёта
    if payload.employee_id is not None:
        employee = db.query(Employee).filter(Employee.id == payload.employee_id).first()
        if not employee:
            raise HTTPException(
                status_code=404,
                detail=f"Сотрудник с ID {payload.employee_id} не найден",
            )
        employees = [employee]
    else:
        employees = db.query(Employee).all()

    total_days = 0
    records_created = 0
    records_updated = 0
    needs_review_count = 0

    for employee in employees:
        schedule = employee.schedule  # может быть None

        for day in daterange(date_from, date_to):
            total_days += 1

            # ПРАВИЛЬНОЕ МЕСТО: внутри цикла по дням!
            lunch_minutes = schedule_lunch_minutes(schedule, day)
            default_hours = schedule_norm_hours(schedule, day)

            day_start = datetime.combine(day, time.min)
            day_end = datetime.combine(day, time.max)

            events = (
                db.query(TurnstileEvent)
                .filter(
                    TurnstileEvent.employee_id == employee.id,
                    TurnstileEvent.datetime >= day_start,
                    TurnstileEvent.datetime <= day_end,
                )
                .order_by(TurnstileEvent.datetime)
                .all()
            )

            if not events:
                # Нет событий за день — запись не создаётся
                continue

            first_in, last_out, fact_hours, needs_review, review_reason = calculate_day_record(
                events, lunch_minutes
            )

            existing = (
                db.query(TimesheetRecord)
                .filter(
                    TimesheetRecord.employee_id == employee.id,
                    TimesheetRecord.date == day,
                )
                .first()
            )

            # Плановые часы заполняются вручную начальником и не должны
            # сбрасываться при пересчёте по данным проходной.
            planned_hours = existing.planned_hours if existing else None

            if planned_hours is not None:
                overtime = max(0.0, round(fact_hours - planned_hours, 2))
            else:
                # Предварительный расчёт сверхурочных по норме из графика
                overtime = max(0.0, round(fact_hours - default_hours, 2))

            if existing:
                existing.first_in = first_in
                existing.last_out = last_out
                existing.fact_hours = fact_hours
                existing.default_hours = default_hours
                existing.overtime = overtime
                existing.lunch_minutes = lunch_minutes
                existing.needs_review = needs_review
                existing.review_reason = review_reason
                records_updated += 1
                if needs_review:
                    needs_review_count += 1
            else:
                new_record = TimesheetRecord(
                    employee_id=employee.id,
                    date=day,
                    first_in=first_in,
                    last_out=last_out,
                    fact_hours=fact_hours,
                    planned_hours=None,
                    default_hours=default_hours,
                    overtime=overtime,
                    lunch_minutes=lunch_minutes,
                    needs_review=needs_review,
                    review_reason=review_reason,
                )
                db.add(new_record)
                records_created += 1
                if needs_review:
                    needs_review_count += 1

    db.commit()

    return TimesheetCalculateResponse(
        total_days=total_days,
        employees_processed=len(employees),
        records_created=records_created,
        records_updated=records_updated,
        needs_review_count=needs_review_count,
    )


@router.get("", response_model=List[TimesheetRecordResponse])
def get_timesheet(
    employee_id: Optional[int] = Query(None, description="ID сотрудника"),
    date_from: Optional[str] = Query(None, description="Дата начала периода YYYY-MM-DD"),
    date_to: Optional[str] = Query(None, description="Дата конца периода YYYY-MM-DD"),
    needs_review: Optional[bool] = Query(None, description="Только записи, требующие проверки"),
    db: Session = Depends(get_db),
):
    """Получить список записей табеля с фильтрацией."""

    query = db.query(TimesheetRecord)

    if employee_id is not None:
        query = query.filter(TimesheetRecord.employee_id == employee_id)

    if date_from:
        query = query.filter(TimesheetRecord.date >= parse_date(date_from, "date_from"))

    if date_to:
        query = query.filter(TimesheetRecord.date <= parse_date(date_to, "date_to"))

    if needs_review is not None:
        query = query.filter(TimesheetRecord.needs_review == needs_review)

    records = query.order_by(TimesheetRecord.date, TimesheetRecord.employee_id).all()
    return records


@router.patch("/{record_id}", response_model=TimesheetRecordResponse)
def update_timesheet(
    record_id: int,
    payload: TimesheetUpdateRequest,
    db: Session = Depends(get_db),
):
    """Ручное заполнение плановых часов и снятие пометки о проверке."""

    record = db.query(TimesheetRecord).filter(TimesheetRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail=f"Запись табеля с ID {record_id} не найдена")

    if payload.planned_hours is not None:
        record.planned_hours = payload.planned_hours
        record.overtime = max(0.0, round(record.fact_hours - record.planned_hours, 2))

    if payload.needs_review is not None:
        record.needs_review = payload.needs_review

    if payload.review_reason is not None:
        record.review_reason = payload.review_reason

    db.commit()
    db.refresh(record)

    return record

# ============================================================================
# ЭНДПОИНТЫ ДЛЯ ОТЧЕТА ТАБЕЛЯ
# ============================================================================

from datetime import date
from calendar import monthrange
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


@router.get("/report")
def get_timesheet_report(
    month: int = Query(..., ge=1, le=12, description="Месяц (1-12)"),
    year: int = Query(..., ge=2020, description="Год"),
    department_id: Optional[int] = Query(None, description="ID подразделения"),
    db: Session = Depends(get_db)
):
    """Получить отчет табеля за месяц"""
    
    _, days_in_month = monthrange(year, month)
    
    # Получаем сотрудников
    query = db.query(Employee)
    if department_id:
        all_dept_ids = get_all_department_ids(department_id, db)
        query = query.filter(Employee.department_id.in_(all_dept_ids))
    
    employees = query.order_by(Employee.full_name).all()
    
    employee_reports = []
    
    for idx, emp in enumerate(employees, start=1):
        # Получаем записи табеля за месяц
        month_start = date(year, month, 1)
        if month == 12:
            month_end = date(year + 1, 1, 1)
        else:
            month_end = date(year, month + 1, 1)
        
        records = db.query(TimesheetRecord).filter(
            TimesheetRecord.employee_id == emp.id,
            TimesheetRecord.date >= month_start,
            TimesheetRecord.date < month_end
        ).all()
        
        records_by_day = {record.date.day: record for record in records}
        
        days_data = {}
        total_hours = 0.0
        
        for day in range(1, days_in_month + 1):
            record = records_by_day.get(day)
            
            if record:
                if record.needs_review:
                    value = "О6"
                elif record.overtime and record.overtime > 0:
                    # Формат: "10.25 (2.25с)" — фактические часы + сверхурочные в скобках
                    fact = round(record.fact_hours, 2)
                    overtime = round(record.overtime, 2)
                    value = f"{fact} ({overtime}с)"
                elif record.fact_hours > 0:
                    value = str(round(record.fact_hours, 2))
                else:
                    value = "в"
                
                days_data[str(day)] = {
                    "value": value,
                    "hours": record.fact_hours,
                    "needs_review": record.needs_review
                }
                
                if not record.needs_review:
                    total_hours += record.fact_hours
            else:
                days_data[str(day)] = {
                    "value": "в",
                    "hours": 0.0,
                    "needs_review": False
                }
        
        # Получаем название подразделения
        dept_name = "-"
        if emp.department_id:
            dept = db.query(Department).filter(Department.id == emp.department_id).first()
            if dept:
                dept_name = dept.name
        
        employee_reports.append({
            "index": idx,
            "full_name": emp.full_name,
            "department": dept_name,
            "days": days_data,
            "total_hours": round(total_hours, 2)
        })
    
    return {
        "month": month,
        "year": year,
        "days_in_month": days_in_month,
        "employees": employee_reports
    }


@router.get("/report/excel")
def get_timesheet_report_excel(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2020),
    department_id: Optional[int] = Query(None),
    db: Session = Depends(get_db)
):
    """Сгенерировать Excel файл с табелем"""
    
    _, days_in_month = monthrange(year, month)
    
    query = db.query(Employee)
    if department_id:
        all_dept_ids = get_all_department_ids(department_id, db)
        query = query.filter(Employee.department_id.in_(all_dept_ids))
    
    employees = query.order_by(Employee.full_name).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"Табель {month}-{year}"
    
    month_names = {
        1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
        5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
        9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь"
    }
    
    # Заголовок
    ws.merge_cells('A1:C1')
    ws.cell(row=1, column=1, value=f"Табель за {month_names[month]} {year}г.")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    
    # Шапка
    ws.cell(row=3, column=1, value="№ п/п")
    ws.cell(row=3, column=2, value="Ф.И.О.")
    ws.cell(row=3, column=3, value="Подразделение")
    
    for day in range(1, days_in_month + 1):
        ws.cell(row=3, column=3 + day, value=str(day))
        ws.column_dimensions[get_column_letter(3 + day)].width = 5
    
    ws.cell(row=3, column=4 + days_in_month, value="Итого")
    ws.column_dimensions[get_column_letter(4 + days_in_month)].width = 10
    
    # Данные
    for idx, emp in enumerate(employees, start=1):
        ws.cell(row=4 + idx, column=1, value=idx)
        ws.cell(row=4 + idx, column=2, value=emp.full_name)
        
        dept_name = "-"
        if emp.department_id:
            dept = db.query(Department).filter(Department.id == emp.department_id).first()
            if dept:
                dept_name = dept.name
        ws.cell(row=4 + idx, column=3, value=dept_name)
        
        month_start = date(year, month, 1)
        if month == 12:
            month_end = date(year + 1, 1, 1)
        else:
            month_end = date(year, month + 1, 1)
        
        records = db.query(TimesheetRecord).filter(
            TimesheetRecord.employee_id == emp.id,
            TimesheetRecord.date >= month_start,
            TimesheetRecord.date < month_end
        ).all()
        
        records_by_day = {record.date.day: record for record in records}
        total_hours = 0.0
        
        for day in range(1, days_in_month + 1):
            record = records_by_day.get(day)
            col = 3 + day
            
            if record:
                if record.needs_review:
                    value = "О6"
                    fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
                elif record.overtime and record.overtime > 0:
                    # Формат: "10.25 (2.25с)"
                    fact = round(record.fact_hours, 2)
                    overtime = round(record.overtime, 2)
                    value = f"{fact} ({overtime}с)"
                    fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
                elif record.fact_hours > 0:
                    value = str(round(record.fact_hours, 2))
                    fill = None
                else:
                    value = "в"
                    fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
                
                ws.cell(row=4 + idx, column=col, value=value)
                if fill:
                    ws.cell(row=4 + idx, column=col).fill = fill
                
                if not record.needs_review:
                    total_hours += record.fact_hours
            else:
                cell = ws.cell(row=4 + idx, column=col, value="в")
                cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
        
        ws.cell(row=4 + idx, column=4 + days_in_month, value=round(total_hours, 2))
        ws.cell(row=4 + idx, column=4 + days_in_month).font = Font(bold=True)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"табель_{month_names[month]}_{year}.xlsx"
    filename_encoded = quote(filename)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{filename_encoded}"
        }
    )