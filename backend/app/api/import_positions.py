from fastapi import APIRouter, UploadFile, File, Depends, HTTPException
from sqlalchemy.orm import Session
from io import BytesIO
from openpyxl import load_workbook

from app.core.security import require_admin
from app.database import get_db
from app.models.employee import Employee

router = APIRouter(prefix="/import", tags=["Import"])


def _norm_tab_number(value) -> str | None:
    """Нормализует табельный номер: убирает '.0' у чисел из Excel."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    # Excel может вернуть число как 12345.0 — приводим к целому виду
    if text.endswith(".0"):
        try:
            text = str(int(float(text)))
        except ValueError:
            pass
    return text


@router.post("/update-positions")
async def update_employee_positions(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user=Depends(require_admin),
):
    """
    Обновить должности существующих сотрудников из Excel файла.
    Файл должен иметь колонки: ФИО, Таб.номер, Подразделение, Должность.
    Новые сотрудники НЕ создаются — обновляется только поле position.
    Поиск: сначала по табельному номеру (приоритет), затем по ФИО.
    """
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=400,
            detail="Неверный формат файла. Требуется Excel (.xlsx или .xls)",
        )

    try:
        content = await file.read()
        wb = load_workbook(filename=BytesIO(content), data_only=True)
        sheet = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка чтения файла: {str(e)}")

    updated_count = 0
    unchanged_count = 0
    not_found_count = 0
    not_found_list = []

    try:
        rows = list(sheet.iter_rows(min_row=2, values_only=True))

        for row_idx, row in enumerate(rows, start=2):
            if not row or not any(cell is not None and str(cell).strip() for cell in row):
                continue

            full_name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else None
            tab_number = _norm_tab_number(row[1]) if len(row) > 1 else None
            position = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None

            if position in ("", "nan"):
                position = None

            # Нет новой должности — нечего обновлять
            if not position:
                continue

            # Ищем сотрудника по табельному номеру (приоритет)
            employee = None
            if tab_number:
                employee = db.query(Employee).filter(Employee.tab_number == tab_number).first()

            # Если не нашли — ищем по ФИО
            if not employee and full_name and full_name.lower() != "nan":
                employee = db.query(Employee).filter(Employee.full_name == full_name).first()

            if employee:
                if employee.position != position:  # Обновляем только если изменилось
                    employee.position = position
                    updated_count += 1
                else:
                    unchanged_count += 1
            else:
                not_found_count += 1
                if len(not_found_list) < 20:  # Сохраняем первые 20
                    not_found_list.append(
                        f"Стр. {row_idx}: {full_name or '—'} (Таб. {tab_number or '—'})"
                    )

        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

    return {
        "success": True,
        "updated": updated_count,
        "unchanged": unchanged_count,
        "not_found": not_found_count,
        "not_found_list": not_found_list,
    }
